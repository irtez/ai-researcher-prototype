import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from openpyxl import Workbook

from file_processing import CsvTable, DocxTable, ExcelTable, read_data_file
from file_tools import TOOL_REGISTRY, tool_read_data_file


class FileProcessingTests(unittest.TestCase):
    def test_csv_returns_raw_text_as_single_table(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            data_dir = Path(tmpdir) / "data"
            data_dir.mkdir()
            csv_path = data_dir / "a.csv"
            text = "compound,value\nAspirin,42\n"
            csv_path.write_text(text, encoding="utf-8")

            tables = read_data_file("a.csv", data_dir=data_dir)

            self.assertEqual(tables, [CsvTable(content=text)])
            self.assertEqual(str(tables[0]), text)

    def test_excel_splits_disconnected_tables_on_same_sheet(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            data_dir = Path(tmpdir) / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "tables.xlsx"
            workbook = Workbook()
            worksheet = workbook.active

            self._write_table(
                worksheet,
                start_row=2,
                start_column=2,
                rows=[
                    ["YSZ - 0%", ""],
                    ["T", "log(Z')"],
                    [248, 6.301029],
                    [295, 5.439332],
                ],
            )
            self._write_table(
                worksheet,
                start_row=2,
                start_column=5,
                rows=[
                    ["YSZ - 0,25%", ""],
                    ["T", "log(Z')"],
                    [199, 7.301029],
                    [246, 6.342422],
                ],
            )
            self._write_table(
                worksheet,
                start_row=2,
                start_column=8,
                rows=[
                    ["YSZ - 2,5%", ""],
                    ["T", "log(3Z')"],
                    [242, 6.270534],
                    [291, 5.409595],
                ],
            )
            workbook.save(workbook_path)

            tables = read_data_file("tables.xlsx", data_dir=data_dir)

            self.assertEqual(len(tables), 3)
            self.assertIsInstance(tables[0], ExcelTable)
            self.assertEqual(
                tables[0].content,
                "\n".join(
                    [
                        "YSZ - 0%,",
                        "T,log(Z')",
                        "248,6.301029",
                        "295,5.439332",
                    ]
                ),
            )
            self.assertEqual(tables[0].sheet_index, 1)
            self.assertEqual(tables[0].sheet_name, "Sheet")
            self.assertEqual(tables[0].top_left, "B2")
            self.assertEqual(tables[0].bottom_right, "C5")
            self.assertIn("range B2:C5", str(tables[0]))
            self.assertIn("YSZ - 0,25%", tables[1].content)
            self.assertEqual(tables[1].top_left, "E2")
            self.assertEqual(tables[1].bottom_right, "F5")
            self.assertIn("YSZ - 2,5%", tables[2].content)
            self.assertEqual(tables[2].top_left, "H2")
            self.assertEqual(tables[2].bottom_right, "I5")

    def test_excel_returns_only_detected_tables_with_sheet_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            data_dir = Path(tmpdir) / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "multi_sheet.xlsx"
            workbook = Workbook()
            first_sheet = workbook.active
            workbook.create_sheet("empty")
            third_sheet = workbook.create_sheet("third")

            self._write_table(
                first_sheet,
                start_row=1,
                start_column=1,
                rows=[["name", "value"], ["A", 1]],
            )
            self._write_table(
                third_sheet,
                start_row=1,
                start_column=1,
                rows=[["name", "value"], ["B", 2]],
            )
            workbook.save(workbook_path)

            tables = read_data_file("multi_sheet.xlsx", data_dir=data_dir)

            self.assertEqual(len(tables), 2)
            self.assertEqual([table.sheet_index for table in tables], [1, 3])
            self.assertEqual([table.sheet_name for table in tables], ["Sheet", "third"])
            self.assertIn("A,1", tables[0].content)
            self.assertIn("B,2", tables[1].content)

    def test_docx_returns_each_word_table_as_text_block(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            data_dir = Path(tmpdir) / "data"
            data_dir.mkdir()
            document_path = data_dir / "tables.docx"
            document = Document()

            first_table = document.add_table(rows=2, cols=2)
            first_table.cell(0, 0).text = "compound"
            first_table.cell(0, 1).text = "value"
            first_table.cell(1, 0).text = "Aspirin"
            first_table.cell(1, 1).text = "42"

            document.add_paragraph("Text between tables.")

            second_table = document.add_table(rows=2, cols=3)
            second_table.cell(0, 0).text = "T"
            second_table.cell(0, 1).text = "log(Z')"
            second_table.cell(0, 2).text = "note"
            second_table.cell(1, 0).text = "248"
            second_table.cell(1, 1).text = "6.301029"
            second_table.cell(1, 2).text = "ok"
            document.save(document_path)

            tables = read_data_file("tables.docx", data_dir=data_dir)

            self.assertEqual(
                tables,
                [
                    DocxTable(content="compound,value\nAspirin,42", index=1),
                    DocxTable(content="T,log(Z'),note\n248,6.301029,ok", index=2),
                ],
            )
            self.assertEqual(str(tables[1]), "Table 2\nT,log(Z'),note\n248,6.301029,ok")

    def test_file_name_must_stay_inside_data_dir(self) -> None:
        with tempfile.TemporaryDirectory() as tmpdir:
            data_dir = Path(tmpdir) / "data"
            data_dir.mkdir()

            with self.assertRaisesRegex(ValueError, "directory components"):
                read_data_file("../a.csv", data_dir=data_dir)

    def test_tool_wrapper_returns_tables_and_errors(self) -> None:
        with patch.dict(os.environ, {"WORKSPACE_DIR": "/tmp"}):
            with patch(
                "file_tools.read_data_file",
                return_value=[CsvTable(content="name,value\nA,1\n")],
            ) as read_file:
                payload = tool_read_data_file("a.csv", "session-a")

        read_file.assert_called_once_with("a.csv", data_dir=Path("/tmp/session-a/inputs"))
        self.assertIn("read_data_file", TOOL_REGISTRY)
        self.assertTrue(payload["ok"])
        self.assertEqual(payload["tables"], ["name,value\nA,1\n"])

        with patch.dict(os.environ, {"WORKSPACE_DIR": "/tmp"}):
            with patch(
                "file_tools.read_data_file",
                return_value=[
                    ExcelTable(
                        content="first",
                        sheet_index=1,
                        sheet_name="alpha",
                        top_left="A1",
                        bottom_right="A1",
                    ),
                    ExcelTable(
                        content="third",
                        sheet_index=3,
                        sheet_name="gamma",
                        top_left="B2",
                        bottom_right="B2",
                    ),
                ],
            ):
                multi_sheet_payload = tool_read_data_file("tables.xlsx", "session-a")

        self.assertTrue(multi_sheet_payload["ok"])
        self.assertEqual(
            multi_sheet_payload["tables"],
            [
                "(sheet 1: alpha, range A1:A1)\nfirst",
                "(sheet 3: gamma, range B2:B2)\nthird",
            ],
        )

        with patch.dict(os.environ, {"WORKSPACE_DIR": "/tmp"}):
            with patch("file_tools.read_data_file", return_value=[]):
                empty_payload = tool_read_data_file("empty.xlsx", "session-a")

        self.assertTrue(empty_payload["ok"])
        self.assertEqual(empty_payload["tables"], [])

        with patch.dict(os.environ, {"WORKSPACE_DIR": "/tmp"}):
            with patch("file_tools.read_data_file", side_effect=FileNotFoundError("missing")):
                error_payload = tool_read_data_file("missing.csv", "session-a")

        self.assertFalse(error_payload["ok"])
        self.assertEqual(error_payload["error_type"], "FileNotFoundError")

    @staticmethod
    def _write_table(worksheet, *, start_row: int, start_column: int, rows: list[list[object]]) -> None:
        for row_offset, row in enumerate(rows):
            for column_offset, value in enumerate(row):
                worksheet.cell(
                    row=start_row + row_offset,
                    column=start_column + column_offset,
                    value=value,
                )


if __name__ == "__main__":
    unittest.main()
